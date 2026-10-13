"""meeting_doc: turn meeting emails (or meetings the agent already structured) into a checked Word document
(meetings.py). week_only=true just reports the week's date range, so the agent can prepare while it still waits
for the user to pick an account."""
from . import tool
from .. import action_results, meetings


@tool("meeting_doc", "Make a Word document of the user's meetings for a week from emails you already read. Pass "
      "messages=[{id, thread_id, subject, from, date, body, ics}] (ics = the invitation's iCalendar text when the "
      "email has one) and/or meetings=[{title, start, end, location, link, organizer, attendees, agenda, prep, "
      "source}] (ISO times). Follows invitation updates and cancellations, flags conflicts, never invents times, "
      "never overwrites a file, then reopens the document and counts the meetings. week: 'this', 'next', 'last' or "
      "a date in the week. week_only=true only returns the week's date range.", "files_write",
      {"messages": {"type": "array", "items": {"type": "object"}},
       "meetings": {"type": "array", "items": {"type": "object"}},
       "week": {"type": "string"}, "week_start": {"type": "string", "enum": ["monday", "sunday"]},
       "timezone": {"type": "string", "description": "IANA zone, e.g. America/New_York (default: this computer's)"},
       "path": {"type": "string", "description": "Where to save the .docx (default: the working folder)"},
       "title": {"type": "string"}, "week_only": {"type": "boolean"}},
      summary=lambda a: "week range" if a.get("week_only") else (a.get("path") or f"{len(a.get('messages') or [])} emails"))
def meeting_doc(ctx, messages=None, meetings_=None, week="this", week_start="monday", timezone="", path="",
                title="", week_only=False, **kw):
    recs = kw.get("meetings", meetings_)
    try:
        tz = meetings.local_zone(timezone or None)
    except Exception:
        return {"text": f"Unknown time zone '{timezone}'. Use an IANA name like America/New_York.", "error": True}
    tz_name = timezone or getattr(tz, "key", None) or str(tz)
    try:
        start, end = meetings.week_range(week or "this", tz, week_start or "monday")
    except ValueError:
        return {"text": f"'{week}' is not a week: use this, next, last or a date like 2026-10-12.", "error": True}
    rng = meetings.range_text(start, end, tz_name)
    if week_only:
        return f"Week: {rng} (from {start.isoformat()} up to {end.isoformat()})."
    msgs = [m for m in (messages or []) if isinstance(m, dict)]
    data = meetings.collect(msgs, [m for m in (recs or []) if isinstance(m, dict)], start, end, tz)
    if not data["active"] and not data["cancelled"] and not data["mentions"]:
        return {"text": f"No meetings found for {rng} in the {len(msgs)} email(s) given. No document was made.",
                "error": False}
    target = ctx.path(path or f"Meetings {start.strftime('%Y-%m-%d')}.docx")
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    target = meetings.free_path(target)
    meetings.write_docx(target, data, start, end, tz_name, title or None, len(msgs))
    excluded = [m.get("subject") for m in msgs if m.get("id") in set(data["skipped"])]
    ok, det = meetings.verify_docx(target, data, excluded)
    lines = [f"Saved {target}", f"Week: {rng}",
             f"Meetings: {len(data['active'])} active, {len(data['cancelled'])} cancelled, "
             f"{len(data['conflicts'])} conflicts, {len(data['mentions'])} mentioned without calendar details",
             f"Emails not about meetings (left out): {len(data['skipped'])}"]
    for e in data["active"]:
        lines.append(f"- {meetings._when(e, start.tzinfo)}: {e['title']}")
    env = action_results.make("file", verified=ok, method=f"reopened the .docx: {det['meetings_in_document']} of "
                              f"{det['expected']} meetings found" + (", unrelated email text found!"
                                                                    if det["unrelated_found"] else ""),
                              changed={"path": str(target)})
    return action_results.attach({"text": "\n".join(lines)}, env)


# ---- write_docx (1.3): a Word document from Markdown ---------------------------------------------------------------
# Found while testing every feature: asked for a Word document, the model had no way to make one (no skill, and the
# system Python it would script with has no python-docx). Aero's own Python has it; Markdown is what models write well.

import re as _re
from pathlib import Path as _Path

_INLINE = _re.compile(r"(\*\*[^*]+\*\*|(?<!\w)__[^_]+__(?!\w)|\*[^*\s][^*]*\*|(?<!\w)_[^_\s][^_]*_(?!\w)|`[^`]+`)")


def _runs(par, text):
    """Markdown inline marks -> runs: **bold**, *italic*, `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part[:2] in ("**", "__") and part[-2:] == part[:2] and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part[0] == "`" and part[-1] == "`" and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part)


def _table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(_re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]   # the --- line
    if not cells:
        return
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(n):
            par = t.cell(i, j).paragraphs[0]
            _runs(par, r[j] if j < len(r) else "")
            if i == 0:
                for run in par.runs:
                    run.bold = True


def markdown_docx(markdown, title=""):
    """A python-docx Document from Markdown: headings, paragraphs, bullet and numbered lists, pipe tables, code
    blocks, quotes, horizontal rules."""
    import docx
    doc = docx.Document()
    if title:
        doc.add_heading(title, 0)
    lines = str(markdown or "").replace("\r\n", "\n").split("\n")
    i, para = 0, []

    def flush():
        if para:
            _runs(doc.add_paragraph(), " ".join(x.strip() for x in para))
            para.clear()
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
        elif not s:
            flush()
        elif m := _re.match(r"(#{1,6})\s+(.*)", s):
            flush()
            doc.add_heading(m.group(2).strip().rstrip("#").strip(), min(len(m.group(1)), 4))
        elif _re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            flush()
            doc.add_paragraph("_" * 40)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows)
            continue
        elif m := _re.match(r"(\s*)[-*+]\s+(.*)", ln):
            flush()
            _runs(doc.add_paragraph(style="List Bullet" if len(m.group(1)) < 2 else "List Bullet 2"), m.group(2))
        elif m := _re.match(r"(\s*)\d+[.)]\s+(.*)", ln):
            flush()
            _runs(doc.add_paragraph(style="List Number" if len(m.group(1)) < 2 else "List Number 2"), m.group(2))
        elif s.startswith(">"):
            flush()
            _runs(doc.add_paragraph(style="Quote"), s.lstrip("> ").strip())
        else:
            para.append(s)
        i += 1
    flush()
    return doc


@tool("write_docx", "Create a Word document (.docx) from Markdown: # headings, paragraphs, - bullets, 1. numbered "
      "lists, | pipe | tables |, **bold**, *italic*, `code`, ``` code blocks, > quotes. Then it reopens the file and "
      "reports what is in it. Use it whenever the user wants a Word file (report, letter, notes, summary).",
      "files_write",
      {"path": {"type": "string", "description": "where to save it, ending in .docx"},
       "markdown": {"type": "string", "description": "the document's content in Markdown"},
       "title": {"type": "string", "description": "optional title at the top"},
       "overwrite": {"type": "boolean", "description": "replace an existing file (default false)"}},
      ["path", "markdown"], summary=lambda a: a.get("path", ""))
def write_docx(ctx, path, markdown, title="", overwrite=False):
    p = ctx.path(path)
    if p.suffix.lower() != ".docx":
        p = p.with_name(p.name + ".docx")
    if p.exists() and not overwrite:
        return {"text": f"{p} already exists. Choose another name, or set overwrite=true to replace it.", "error": True}
    p.parent.mkdir(parents=True, exist_ok=True)
    markdown_docx(markdown, title).save(str(p))
    import docx
    d = docx.Document(str(p))                      # read it back: what really is in the file
    heads = [x.text for x in d.paragraphs if x.style.name.startswith(("Heading", "Title")) and x.text.strip()]
    words = sum(len(x.text.split()) for x in d.paragraphs) + sum(
        len(c.text.split()) for t in d.tables for row in t.rows for c in row.cells)
    text = (f"Wrote {p} ({_Path(p).stat().st_size:,} bytes): {words:,} words, {len(d.paragraphs)} paragraphs, "
            f"{len(d.tables)} tables" + (f"; headings: {', '.join(heads[:12])}" if heads else "") + ".")
    env = action_results.make("file", verified=True, method="reopened the .docx", changed={"path": str(p)})
    return action_results.attach({"text": text}, env)
