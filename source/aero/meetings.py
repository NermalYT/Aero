"""Meetings from email, as a Word document: the deterministic half of "check my Gmail for meetings this week and make
me a document with everything I need".

The agent gets the messages (a Gmail MCP server, Outlook, or Aero's browser after the user signs in) and passes them
here as plain data. This module never reads a mailbox itself. It:

  1. works out the week in the user's own time zone (Monday to Sunday, or Sunday to Saturday), with the range
     written out in the document;
  2. reads calendar invitations (iCalendar VEVENTs: times with their zones, title, place, organizer, attendees, UID,
     SEQUENCE, STATUS, METHOD:CANCEL) and follows updates: the highest SEQUENCE of a UID wins, cancellations remove a
     meeting from the active list and show it as cancelled;
  3. keeps meetings the agent already structured, after checking their dates;
  4. never invents details: an email that mentions a meeting but carries no calendar data is listed as such, with
     the sentence that mentions it, and no made-up time;
  5. marks overlapping meetings as conflicts;
  6. writes a .docx (python-docx) next to nothing it would overwrite, then reopens it and counts the meetings in it.
"""
import datetime as dt
import re
from pathlib import Path

try:
    from zoneinfo import ZoneInfo
except ImportError:            # pragma: no cover
    ZoneInfo = None

VIDEO = re.compile(r"https?://[^\s<>\"']*(zoom\.us/[jw]/|meet\.google\.com/|teams\.microsoft\.com/l/meetup-join|"
                   r"teams\.live\.com/meet|webex\.com/|whereby\.com/|gotomeeting\.com/)[^\s<>\"')\]]*", re.I)
MEETING_WORDS = re.compile(r"\b(meeting|meet|call|sync|stand-?up|interview|appointment|1:1|one-on-one|webinar|demo|"
                           r"review session|catch[- ]up)\b", re.I)


# ------------------------------------------------------------------------------------------------ the week

def local_zone(name=None):
    if name and ZoneInfo:
        return ZoneInfo(name)
    return dt.datetime.now().astimezone().tzinfo


def week_range(when="this", tz=None, week_start="monday", now=None):
    """(start, end) aware datetimes of a week in tz. when: 'this', 'next', 'last' or a date (YYYY-MM-DD) inside the
    week. end is exclusive (the next week's start)."""
    tz = tz or local_zone()
    now = (now or dt.datetime.now(tz)).astimezone(tz)
    if when in ("this", "", None):
        day = now.date()
    elif when == "next":
        day = now.date() + dt.timedelta(days=7)
    elif when == "last":
        day = now.date() - dt.timedelta(days=7)
    else:
        day = dt.date.fromisoformat(str(when)[:10])
    first = 0 if week_start == "monday" else 6           # date.weekday(): Monday 0 ... Sunday 6
    back = (day.weekday() - first) % 7
    start_day = day - dt.timedelta(days=back)
    start = dt.datetime.combine(start_day, dt.time(0, 0), tzinfo=tz)
    end = dt.datetime.combine(start_day + dt.timedelta(days=7), dt.time(0, 0), tzinfo=tz)
    return start, end


def range_text(start, end, tz_name=""):
    last = end - dt.timedelta(days=1)
    same_year = start.year == last.year
    a = start.strftime("%a %b %d") + ("" if same_year else start.strftime(", %Y"))
    b = last.strftime("%a %b %d, %Y")
    return f"{a} – {b}" + (f" ({tz_name})" if tz_name else "")


# ------------------------------------------------------------------------------------------------ iCalendar

def _unfold(text):
    return re.sub(r"\r?\n[ \t]", "", text or "")


def _ics_value(line):
    """('DTSTART', {'TZID': 'Europe/Paris'}, '20261013T140000')"""
    head, _, value = line.partition(":")
    parts = head.split(";")
    params = {}
    for p in parts[1:]:
        k, _, v = p.partition("=")
        params[k.upper()] = v.strip('"')
    return parts[0].upper(), params, value


def _ics_time(value, params, default_tz):
    v = value.strip()
    try:
        if params.get("VALUE") == "DATE" or re.fullmatch(r"\d{8}", v):
            d = dt.date(int(v[:4]), int(v[4:6]), int(v[6:8]))
            return dt.datetime.combine(d, dt.time(0, 0), tzinfo=default_tz), True
        t = dt.datetime.strptime(v.rstrip("Z"), "%Y%m%dT%H%M%S")
        if v.endswith("Z"):
            return t.replace(tzinfo=dt.timezone.utc), False
        tzid = params.get("TZID")
        tz = None
        if tzid and ZoneInfo:
            try:
                tz = ZoneInfo(tzid)
            except Exception:
                tz = _WINDOWS_TZ.get(tzid)
        return t.replace(tzinfo=tz or default_tz), False
    except (ValueError, IndexError):
        return None, False


# Outlook writes Windows zone names into TZID
_WINDOWS_TZ = {}
if ZoneInfo:
    for _w, _i in (("Eastern Standard Time", "America/New_York"), ("Central Standard Time", "America/Chicago"),
                   ("Mountain Standard Time", "America/Denver"), ("Pacific Standard Time", "America/Los_Angeles"),
                   ("GMT Standard Time", "Europe/London"), ("W. Europe Standard Time", "Europe/Berlin"),
                   ("Romance Standard Time", "Europe/Paris"), ("India Standard Time", "Asia/Kolkata"),
                   ("Tokyo Standard Time", "Asia/Tokyo"), ("AUS Eastern Standard Time", "Australia/Sydney"),
                   ("UTC", "UTC")):
        try:
            _WINDOWS_TZ[_w] = ZoneInfo(_i)
        except Exception:
            pass


def parse_ics(text, default_tz=None):
    """VEVENTs in an iCalendar text: [{uid, sequence, title, start, end, all_day, location, organizer, attendees,
    status, method, description, links}]."""
    default_tz = default_tz or local_zone()
    text = _unfold(text)
    method = ""
    m = re.search(r"^METHOD:(.+)$", text, re.M)
    if m:
        method = m.group(1).strip().upper()
    out = []
    for block in re.findall(r"BEGIN:VEVENT(.*?)END:VEVENT", text, re.S):
        ev = {"uid": "", "sequence": 0, "title": "", "start": None, "end": None, "all_day": False, "location": "",
              "organizer": "", "attendees": [], "status": "", "method": method, "description": "", "links": []}
        for line in block.strip().splitlines():
            if ":" not in line:
                continue
            key, params, value = _ics_value(line)
            value = value.replace("\\n", "\n").replace("\\,", ",").replace("\\;", ";")
            if key == "UID":
                ev["uid"] = value.strip()
            elif key == "SEQUENCE":
                ev["sequence"] = int(re.sub(r"\D", "", value) or 0)
            elif key == "SUMMARY":
                ev["title"] = value.strip()
            elif key == "DTSTART":
                ev["start"], ev["all_day"] = _ics_time(value, params, default_tz)
            elif key == "DTEND":
                ev["end"], _ = _ics_time(value, params, default_tz)
            elif key == "LOCATION":
                ev["location"] = value.strip()
            elif key == "ORGANIZER":
                ev["organizer"] = params.get("CN") or value.replace("mailto:", "").strip()
            elif key == "ATTENDEE":
                ev["attendees"].append(params.get("CN") or value.replace("mailto:", "").strip())
            elif key == "STATUS":
                ev["status"] = value.strip().upper()
            elif key == "DESCRIPTION":
                ev["description"] = value.strip()[:2000]
        ev["links"] = sorted(set(m.group(0) for m in VIDEO.finditer(ev["description"] + " " + ev["location"])))
        if ev["start"] is not None:
            out.append(ev)
    return out


# ------------------------------------------------------------------------------------------------ collecting

def _cancelled_subject(subject):
    return bool(re.match(r"^\s*(canceled|cancelled|abgesagt|annulé)\b", subject or "", re.I))


def collect(messages=None, meetings=None, start=None, end=None, tz=None):
    """Build the meeting list. messages: [{id, thread_id, subject, from, date, body, ics}]. meetings: records the
    agent structured itself ({title, start, end, location, link, organizer, attendees, agenda, prep, source}).
    Returns {"active", "cancelled", "mentions", "conflicts", "skipped"}."""
    tz = tz or local_zone()
    by_uid, loose, mentions, skipped = {}, [], [], []
    for msg in messages or []:
        src = {"id": msg.get("id") or "", "subject": (msg.get("subject") or "")[:200], "from": msg.get("from") or "",
               "date": msg.get("date") or "", "thread": msg.get("thread_id") or ""}
        evs = parse_ics(msg.get("ics") or "", tz) if msg.get("ics") else []
        if not evs:
            body = msg.get("body") or msg.get("snippet") or ""
            if MEETING_WORDS.search(src["subject"] + " " + body):
                sent = next((s.strip() for s in re.split(r"(?<=[.!?])\s+|\n", body) if MEETING_WORDS.search(s)), "")
                mentions.append({"source": src, "sentence": sent[:300],
                                 "links": sorted(set(m.group(0) for m in VIDEO.finditer(body)))[:3]})
            else:
                skipped.append(src["id"])
            continue
        for ev in evs:
            ev = dict(ev, source=src)
            if ev["method"] == "CANCEL" or ev["status"] == "CANCELLED" or _cancelled_subject(src["subject"]):
                ev["status"] = "CANCELLED"
            key = ev["uid"] or f"{ev['title'].lower()}|{ev['start'].isoformat()}"
            prev = by_uid.get(key)
            if prev is None or ev["sequence"] > prev["sequence"] or \
                    (ev["sequence"] == prev["sequence"] and ev["status"] == "CANCELLED"):
                if prev is not None:
                    ev["history"] = (prev.get("history") or []) + [{"sequence": prev["sequence"],
                                                                     "start": prev["start"], "status": prev["status"]}]
                by_uid[key] = ev
    for m in meetings or []:
        try:
            st = _parse_when(m.get("start"), tz)
        except ValueError:
            skipped.append(m.get("title") or "(meeting without a valid start)")
            continue
        en = None
        if m.get("end"):
            try:
                en = _parse_when(m["end"], tz)
            except ValueError:
                en = None
        loose.append({"uid": "", "sequence": 0, "title": str(m.get("title") or "(untitled)")[:200], "start": st,
                      "end": en, "all_day": False, "location": m.get("location") or "",
                      "organizer": m.get("organizer") or "", "attendees": list(m.get("attendees") or [])[:50],
                      "status": str(m.get("status") or "").upper(), "description": m.get("agenda") or "",
                      "links": [m["link"]] if m.get("link") else [], "prep": list(m.get("prep") or [])[:20],
                      "source": {"id": str(m.get("source") or ""), "subject": "", "from": "", "date": ""}})
    events = list(by_uid.values())
    for ev in loose:                                      # agent records that duplicate an invitation are dropped
        if not any(e["title"].lower() == ev["title"].lower() and abs((e["start"] - ev["start"]).total_seconds()) < 60
                   for e in events):
            events.append(ev)
    if start is not None:
        events = [e for e in events if start <= e["start"].astimezone(start.tzinfo) < end]
    events.sort(key=lambda e: e["start"])
    active = [e for e in events if e["status"] != "CANCELLED"]
    cancelled = [e for e in events if e["status"] == "CANCELLED"]
    conflicts = []
    for i, a in enumerate(active):
        a_end = a["end"] or a["start"] + dt.timedelta(minutes=30)
        for b in active[i + 1:]:
            if b["start"] < a_end and not a["all_day"] and not b["all_day"]:
                conflicts.append((a, b))
    return {"active": active, "cancelled": cancelled, "mentions": mentions, "conflicts": conflicts, "skipped": skipped}


def _parse_when(v, tz):
    if isinstance(v, dt.datetime):
        return v if v.tzinfo else v.replace(tzinfo=tz)
    s = str(v or "").strip()
    if not s:
        raise ValueError("empty")
    t = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    return t if t.tzinfo else t.replace(tzinfo=tz)


# ------------------------------------------------------------------------------------------------ document

def _when(e, tz):
    s = e["start"].astimezone(tz)
    if e["all_day"]:
        return s.strftime("%a %b %d") + " (all day)"
    out = s.strftime("%a %b %d, %I:%M %p").replace(" 0", " ")
    if e.get("end"):
        out += " – " + e["end"].astimezone(tz).strftime("%I:%M %p").lstrip("0")
    return out


def write_docx(path, data, start, end, tz_name, title=None, n_messages=0):
    from docx import Document
    doc = Document()
    tz = start.tzinfo
    doc.add_heading(title or f"Meetings, {range_text(start, end, tz_name)}", 0)
    doc.add_paragraph(f"{len(data['active'])} meeting{'s' if len(data['active']) != 1 else ''} this week"
                      + (f", {len(data['cancelled'])} cancelled" if data["cancelled"] else "")
                      + (f", {len(data['conflicts'])} conflict{'s' if len(data['conflicts']) != 1 else ''}"
                         if data["conflicts"] else "") + ".")
    if data["active"]:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else t.style
        for i, h in enumerate(("When", "Meeting", "Where / link", "Organizer")):
            t.rows[0].cells[i].text = h
        for e in data["active"]:
            r = t.add_row().cells
            r[0].text = _when(e, tz)
            r[1].text = e["title"]
            r[2].text = e["location"] or (e["links"][0] if e["links"] else "")
            r[3].text = e["organizer"]
    for e in data["active"]:
        doc.add_heading(e["title"], level=1)
        rows = [("When", _when(e, tz)), ("Where", e["location"]), ("Video link", ", ".join(e["links"])),
                ("Organizer", e["organizer"]), ("Attendees", ", ".join(e["attendees"][:25]) +
                                                 (f" and {len(e['attendees']) - 25} more" if len(e["attendees"]) > 25 else "")),
                ("Agenda / details", e.get("description") or ""), ("Prep", "; ".join(e.get("prep") or []))]
        for k, v in rows:
            if v:
                p = doc.add_paragraph()
                p.add_run(k + ": ").bold = True
                p.add_run(v)
        if e.get("history"):
            doc.add_paragraph("Changed: earlier versions of this invitation were updated (the latest is shown).")
        src = e.get("source") or {}
        if src.get("subject") or src.get("id"):
            doc.add_paragraph(f"Source: email \"{src.get('subject') or src.get('id')}\"" +
                              (f" from {src['from']}" if src.get("from") else "") +
                              (f", {src['date']}" if src.get("date") else ""), style="Intense Quote"
                              if "Intense Quote" in [s.name for s in doc.styles] else None)
    if data["conflicts"]:
        doc.add_heading("Conflicts", level=1)
        for a, b in data["conflicts"]:
            doc.add_paragraph(f"{a['title']} ({_when(a, tz)}) overlaps {b['title']} ({_when(b, tz)}).",
                              style="List Bullet")
    if data["cancelled"]:
        doc.add_heading("Cancelled", level=1)
        for e in data["cancelled"]:
            doc.add_paragraph(f"{e['title']}, was {_when(e, tz)}", style="List Bullet")
    if data["mentions"]:
        doc.add_heading("Mentioned in email without calendar details", level=1)
        doc.add_paragraph("These emails talk about a meeting but carry no invitation, so no time is given here.")
        for m in data["mentions"]:
            s = m["source"]
            doc.add_paragraph(f"\"{s['subject']}\" from {s['from'] or 'unknown sender'}"
                              + (f" ({s['date']})" if s["date"] else "") + (f": {m['sentence']}" if m["sentence"] else ""),
                              style="List Bullet")
    doc.add_paragraph(f"Made by Aero on {dt.datetime.now(tz).strftime('%Y-%m-%d %H:%M')} from {n_messages} email"
                      f"{'s' if n_messages != 1 else ''}. Only details found in those emails are included.")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def free_path(path):
    """path, or 'name (2).docx' ... when it exists: never overwrite a document."""
    p = Path(path)
    if not p.exists():
        return p
    for i in range(2, 200):
        q = p.with_name(f"{p.stem} ({i}){p.suffix}")
        if not q.exists():
            return q
    raise FileExistsError(str(p))


def verify_docx(path, data, excluded_subjects=()):
    """Reopen the document: it must load, hold one heading per active meeting, and none of the emails that were not
    about meetings. Returns (ok, details)."""
    from docx import Document
    doc = Document(str(path))
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    found = sum(1 for e in data["active"] if e["title"] in heads)
    body = "\n".join(p.text for p in doc.paragraphs)
    leaks = [s for s in excluded_subjects if s and s in body]
    ok = found == len(data["active"]) and not leaks
    return ok, {"meetings_in_document": found, "expected": len(data["active"]), "unrelated_found": leaks}
