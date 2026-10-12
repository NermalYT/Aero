"""Finance workspace: anthropics/financial-services (Apache-2.0) built into Aero.

48 skills from the repository's vertical plugins (financial analysis, investment banking, equity research, private
equity, fund admin, operations) are vendored in skills_builtin/packs/finance. They are drafts-for-professional-review
workflows: Aero never gives investment advice and never executes transactions (the full-auto authority model also
denies anything that moves money). The paid data connectors the repository lists are not included.

Aero adds the checks the skills call for, working on the user's files:
  xlsx_audit      formula errors, hard-coded numbers inside formulas, broken row patterns, references to empty cells,
                  circular references (no Excel needed)
  xlsx_recalc     recalculate a workbook with LibreOffice when it is installed (the repository's recalc.py is missing)
  dcf_validate    the dcf-model skill's validate_dcf.py (terminal growth vs WACC, TV share of EV, errors)
  deck_numbers    the ib-check-deck skill's extract_numbers.py on a .pptx (numbers that disagree across slides)
"""
import re
import shutil
import tempfile
from collections import defaultdict
from pathlib import Path

from . import Pack, Requirement, register, toolchain as tc
from .. import action_results, capabilities
from ..tools import tool

SKILLS = Path(__file__).resolve().parent.parent / "skills_builtin" / "packs" / "finance"
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_REF = re.compile(r"(?:'?([^'!]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)")
_NUM_IN_FORMULA = re.compile(r"(?<![A-Z$\d.])(\d+\.\d+|\d{2,})(?![\d.]*[A-Z(])")


def skills_list():
    out = []
    for d in sorted(p for p in SKILLS.iterdir() if (p / "SKILL.md").exists()):
        txt = (d / "SKILL.md").read_text(encoding="utf-8", errors="replace")[:3000]
        m = re.search(r"^description:\s*(.+)$", txt, re.M)
        out.append({"name": d.name, "description": (m.group(1) if m else "").strip()[:300]})
    return out


def skill_hint(text):
    """The finance skills a request names (by words in the skill name), at most four."""
    t = (text or "").lower()
    hits = []
    for s in skills_list():
        words = [w for w in s["name"].split("-") if len(w) > 2 and w not in ("model", "analysis", "builder", "xls")]
        if s["name"].replace("-", " ") in t or (words and all(w in t for w in words)):
            hits.append(s["name"])
    alias = {"dcf": "dcf-model", "lbo": "lbo-model", "comps": "comps-analysis", "comparable": "comps-analysis",
             "3-statement": "3-statement-model", "three statement": "3-statement-model", "merger": "merger-model",
             "ic memo": "ic-memo", "cim": "cim-builder", "teaser": "teaser", "kyc": "kyc-rules", "earnings": "earnings-analysis",
             "reconcil": "gl-recon", "pitch": "pitch-deck", "accrual": "accrual-schedule", "nav": "nav-tieout"}
    for k, v in alias.items():
        if k in t and v not in hits:
            hits.append(v)
    return hits[:4]


def audit(path):
    import openpyxl
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    issues = []
    deps = defaultdict(set)
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        rows = defaultdict(list)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                cached = wv[c.coordinate].value
                if isinstance(cached, str) and cached.strip() in ERRORS:
                    issues.append(("error", ws.title, c.coordinate, f"shows {cached}"))
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                f = v[1:]
                for m in _REF.finditer(f.upper()):
                    sheet = m.group(1) or ws.title
                    ref = f"{m.group(2)}{m.group(3)}"
                    deps[(ws.title, c.coordinate)].add((sheet, ref))
                    if sheet in wb_f.sheetnames and wb_f[sheet][ref].value in (None, ""):
                        issues.append(("warning", ws.title, c.coordinate, f"refers to empty cell {sheet}!{ref}"))
                    elif m.group(1) and sheet not in wb_f.sheetnames:
                        issues.append(("error", ws.title, c.coordinate, f"refers to missing sheet '{sheet}'"))
                consts = [x for x in _NUM_IN_FORMULA.findall(re.sub(r'"[^"]*"', "", f)) if x not in ("100", "12", "365")]
                if consts:
                    issues.append(("warning", ws.title, c.coordinate, f"hard-coded number(s) {', '.join(consts[:3])} "
                                                                      f"inside a formula: {v[:80]}"))
                shape = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "R", f.upper())
                rows[c.row].append((c.column, c.coordinate, shape))
        for r, cells in rows.items():
            shapes = [s for _, _, s in sorted(cells)]
            if len(shapes) >= 4:
                common = max(set(shapes), key=shapes.count)
                for _col, coord, s in sorted(cells):
                    if s != common and shapes.count(common) >= len(shapes) - 1:
                        issues.append(("warning", ws.title, coord, "formula differs from the rest of its row"))
    seen, stack = set(), set()

    def cyc(node, path_):
        if node in stack:
            return path_[path_.index(node):] if node in path_ else [node]
        if node in seen:
            return None
        seen.add(node)
        stack.add(node)
        for d in deps.get(node, ()):
            r = cyc(d, path_ + [node])
            if r:
                return r
        stack.discard(node)
        return None
    for n in list(deps):
        loop = cyc(n, [])
        if loop:
            issues.append(("error", n[0], n[1], "circular reference: " + " → ".join(f"{s}!{c}" for s, c in loop[:6])))
            break
    no_cache = sum(1 for ws in wb_f.worksheets for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("=") and wb_v[ws.title][c.coordinate].value is None)
    return issues, no_cache


@tool("xlsx_audit", "Audit an Excel workbook's formulas without Excel: error values (#REF!, #DIV/0! …), references to "
      "empty cells or missing sheets, hard-coded numbers inside formulas, formulas that break their row's pattern, "
      "circular references.", "integrations", {"path": {"type": "string"}}, ["path"],
      summary=lambda a: a.get("path", ""))
def xlsx_audit(ctx, path):
    p = ctx.path(path)
    if not p.exists():
        return {"text": f"No file at {p}.", "error": True}
    try:
        issues, no_cache = audit(p)
    except Exception as e:  # noqa: BLE001
        return {"text": f"Couldn't read {p.name}: {e}", "error": True}
    lines = [f"{p.name}: {sum(1 for i in issues if i[0] == 'error')} errors, "
             f"{sum(1 for i in issues if i[0] == 'warning')} warnings"]
    if no_cache:
        lines.append(f"{no_cache} formulas have no calculated value saved (written by a script, never opened in a "
                     "spreadsheet app): xlsx_recalc, or open it in Excel, before trusting values.")
    lines += [f"- [{k}] {s}!{c}: {m}" for k, s, c, m in issues[:120]]
    return "\n".join(lines)


def _soffice():
    for name in ("soffice", "libreoffice"):
        p = tc.which(name)
        if p:
            return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(cand).exists():
            return cand
    return None


@tool("xlsx_recalc", "Recalculate every formula in a workbook with LibreOffice (headless) and save the values, so "
      "audits and the user see real numbers. Needs LibreOffice installed.", "dev", {"path": {"type": "string"}},
      ["path"], summary=lambda a: a.get("path", ""))
def xlsx_recalc(ctx, path):
    p = ctx.path(path)
    so = _soffice()
    if not so:
        return {"text": "LibreOffice isn't installed, so Aero can't recalculate. The values are calculated when the "
                        "file is opened in Excel or LibreOffice.", "error": True}
    out = Path(tempfile.mkdtemp(prefix="aero-recalc-"))
    r = tc.run([so, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(p)], timeout=300,
               cancel=ctx.cancel)
    res = out / (p.stem + ".xlsx")
    if r["rc"] != 0 or not res.exists():
        return {"text": f"LibreOffice failed: {(r['err'] or r['out'])[-500:]}", "error": True}
    backup = p.with_name(p.stem + ".before-recalc" + p.suffix)
    shutil.copy2(p, backup)
    shutil.move(str(res), p)
    env = action_results.make("file", verified=p.exists(), method="LibreOffice converted and saved",
                              changed={"path": str(p), "backup": str(backup)})
    return action_results.attach({"text": f"Recalculated {p} (the previous file is {backup.name})."}, env)


def _script(rel, args, timeout=180):
    return tc.run_python(SKILLS / rel, args, timeout)


@tool("dcf_validate", "Check a DCF model workbook with the dcf-model skill's validator: formula errors, terminal growth "
      "below WACC, WACC range, terminal value share of enterprise value.", "integrations",
      {"path": {"type": "string"}}, ["path"], summary=lambda a: a.get("path", ""))
def dcf_validate(ctx, path):
    rc, out, err = _script("dcf-model/scripts/validate_dcf.py", [ctx.path(path)])
    return {"text": (out or err)[:20000], "error": rc != 0 and not out}


@tool("deck_numbers", "Find numbers that disagree across a presentation (the ib-check-deck skill's extract_numbers): "
      "pass a .pptx or a Markdown export of the deck.", "integrations", {"path": {"type": "string"}}, ["path"],
      summary=lambda a: a.get("path", ""))
def deck_numbers(ctx, path):
    p = ctx.path(path)
    src = p
    if p.suffix.lower() == ".pptx":
        from pptx import Presentation
        md = []
        for i, slide in enumerate(Presentation(str(p)).slides, 1):
            md.append(f"## Slide {i}")
            for sh in slide.shapes:
                if sh.has_text_frame:
                    md.append(sh.text_frame.text)
        src = Path(tempfile.mkdtemp(prefix="aero-deck-")) / (p.stem + ".md")
        src.write_text("\n".join(md), encoding="utf-8")
    rc, out, err = _script("ib-check-deck/scripts/extract_numbers.py", [src, "--check"])
    return {"text": (out or err)[:20000], "error": rc != 0 and not out}


for _n in ("xlsx_audit", "dcf_validate", "deck_numbers"):
    capabilities.register(_n, ["finance.check", "spreadsheet.read"], "local_read")
capabilities.register("xlsx_recalc", ["spreadsheet.edit"], "local_write", verify="file replaced, backup kept")


def _ok():
    n = len(skills_list())
    return n > 0, f"{n} skills built in"


register(Pack("financial-services", "Financial services skills", "finance",
              "DCF, LBO, comps, 3-statement, merger models, IC memos, CIMs, earnings notes, GL recon, KYC "
              "(drafts for professional review; never advice or trades)",
              "https://github.com/anthropics/financial-services", "Apache-2.0",
              triggers=[r"\b(dcf|lbo|wacc|comps?|comparable compan(y|ies)|3[- ]statement|three[- ]statement|merger model|"
                        r"accretion|dilution|ic memo|cim|teaser|pitch ?book|earnings (note|preview|analysis)|"
                        r"initiat(e|ing) coverage|gl recon|reconcil\w+|month[- ]end|accruals?|nav tie|kyc|due diligence|"
                        r"valuation|financial model)\b"],
              tools=["use_skill", "xlsx_audit", "xlsx_recalc", "dcf_validate", "deck_numbers", "read_file", "write_file"],
              skills=[], requirements=[Requirement("skills", "financial-services skills", _ok),
                                       Requirement("libreoffice", "LibreOffice (only for xlsx_recalc)",
                                                   lambda: (bool(_soffice()), _soffice() or "not installed"),
                                                   how="install LibreOffice to recalculate workbooks", optional=True)],
              panel=lambda: {"skills": skills_list()}, skill_hint=skill_hint,
              notes="Data connectors from the repository need paid provider accounts and are not included."))
