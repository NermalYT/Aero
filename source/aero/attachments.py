"""Turn any attached file into something the model can read.

Images stay images (sent to vision models via the mmproj projector). Everything else is
converted to text: documents, spreadsheets, slides, archives, code, and for unknown binaries a
printable-strings extract plus a hex header so the model can still reason about the file.
"""
import base64
import io
import json
import mimetypes
import re
import uuid
import zipfile
from pathlib import Path

from .config import UPLOADS

MAX_CHARS = 400_000
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".ico", ".heic"}


def _decode(b: bytes):
    if b.startswith(b"\xff\xfe") or b.startswith(b"\xfe\xff"):
        return b.decode("utf-16", "replace")
    if b.startswith(b"\xef\xbb\xbf"):
        return b[3:].decode("utf-8", "replace")
    try:
        return b.decode("utf-8")
    except UnicodeDecodeError:
        pass
    sample = b[:8192]
    if sample.count(b"\x00") > len(sample) * 0.1:
        return None
    return b.decode("cp1252", "replace")


def _strings(b: bytes, min_len=5, limit=20000):
    out, total = [], 0
    for m in re.finditer(rb"[\x20-\x7e]{%d,}" % min_len, b):
        s = m.group().decode("ascii")
        out.append(s)
        total += len(s) + 1
        if total > limit:
            break
    return "\n".join(out)


def _html_text(s: str):
    s = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h\d>|</tr>", "\n", s)
    s = re.sub(r"<[^>]+>", " ", s)
    import html
    s = html.unescape(s)
    s = re.sub(r"[ \t]+", " ", s)
    return re.sub(r"\n\s*\n+", "\n\n", s).strip()


OCR_PAGES = 60                 # scanned PDFs: pages read with OCR at most (about a second each)


def _page_ocr(p):
    """The text of a scanned page: OCR of the images on it (1.3; Windows OCR, see ocr.py)."""
    from . import ocr
    out = []
    try:
        imgs = list(p.images)[:4]
    except Exception:  # noqa: BLE001
        return ""
    for img in imgs:
        t = ocr.image_text(img.data)
        if t:
            out.append(t)
    return "\n".join(out)


def _pdf(b):
    from pypdf import PdfReader

    from . import ocr
    r = PdfReader(io.BytesIO(b))
    pages, scanned, read = [], 0, 0
    can_ocr = ocr.available()
    for i, p in enumerate(r.pages):
        try:
            t = p.extract_text() or ""
        except Exception:  # noqa: BLE001
            pages.append(f"--- page {i+1} --- (unreadable)")
            continue
        if len(t.strip()) < 40:
            scanned += 1
            if can_ocr and read < OCR_PAGES:
                o = _page_ocr(p)
                read += 1
                if o:
                    pages.append(f"--- page {i+1} (text read with OCR) ---\n{o}")
                    continue
        pages.append(f"--- page {i+1} ---\n{t}")
    txt = "\n".join(pages)
    if scanned and not can_ocr:
        txt += "\n\n[Note: little extractable text; this PDF may be scanned images.]"
    elif scanned > read:
        txt += f"\n\n[Note: {scanned - read} more scanned pages weren't read with OCR (limit {OCR_PAGES}).]"
    return txt


def _docx(b):
    import docx
    d = docx.Document(io.BytesIO(b))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(b):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"### Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 2000:
                out.append("... (truncated after 2000 rows)")
                break
            out.append(",".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _pptx(b):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(b))
    out = []
    for i, s in enumerate(prs.slides, 1):
        out.append(f"--- slide {i} ---")
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
        if s.has_notes_slide:
            out.append("Notes: " + s.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _zip(b):
    z = zipfile.ZipFile(io.BytesIO(b))
    out = ["Archive contents:"]
    budget = 120_000
    for info in z.infolist()[:500]:
        out.append(f"  {info.filename}  ({info.file_size:,} bytes)")
    for info in z.infolist():
        if budget <= 0 or info.is_dir() or info.file_size > 200_000:
            continue
        t = _decode(z.read(info))
        if t:
            out.append(f"\n--- {info.filename} ---\n{t[:budget]}")
            budget -= len(t)
    return "\n".join(out)


AUDIO_EXT = {".wav", ".mp3", ".flac", ".ogg", ".oga"}


def _transcript(name, data, ext):
    """1.3: what is said in an attached recording (whisper.cpp, when voice typing is set up)."""
    try:
        from .integrations import voice
        if not voice.ready():
            return None
        t = voice.transcribe_bytes(data, ext)
    except Exception as e:  # noqa: BLE001
        return f"[{name}: an audio file whisper.cpp couldn't transcribe ({e})]"
    return f"[Transcript of {name} (speech to text on this PC with whisper.cpp)]\n{t or '(no speech found)'}"


def extract(name: str, data: bytes):
    """Returns (kind, text_or_None, note)."""
    ext = Path(name).suffix.lower()
    mime = mimetypes.guess_type(name)[0] or ""
    if ext in IMAGE_EXT or mime.startswith("image/"):
        return "image", None, ""
    try:
        if ext == ".pdf":
            return "file", _pdf(data), ""
        if ext in (".docx", ".docm"):
            return "file", _docx(data), ""
        if ext in (".xlsx", ".xlsm"):
            return "file", _xlsx(data), ""
        if ext == ".pptx":
            return "file", _pptx(data), ""
        if ext in (".zip", ".jar", ".apk", ".nupkg", ".whl"):
            return "file", _zip(data), ""
        if ext in (".html", ".htm", ".xhtml"):
            t = _decode(data) or ""
            return "file", _html_text(t) + "\n\n[raw HTML length: %d chars]" % len(t), ""
        if ext == ".ipynb":
            nb = json.loads(data)
            cells = [f"# [{c.get('cell_type')}]\n" + "".join(c.get("source", [])) for c in nb.get("cells", [])]
            return "file", "\n\n".join(cells), ""
    except Exception as e:  # fall through to generic handling
        note = f"(structured parse failed: {e}; showing raw text)"
    else:
        note = ""
    media = ext in AUDIO_EXT or mime.startswith(("audio/", "video/"))
    if ext in AUDIO_EXT:
        got = _transcript(name, data, ext)
        if got:
            return "file", got, ""
    t = None if media else _decode(data)       # a recording's bytes are not text, whatever they decode as
    if t is not None:
        return "file", t, note
    if media:
        return "file", f"[{mime} file, {len(data):,} bytes. Audio/video content cannot be read directly; "\
                       f"metadata strings follow]\n" + _strings(data[:2_000_000], 6, 4000), ""
    hexhead = " ".join(f"{x:02x}" for x in data[:64])
    return "file", (f"[binary file, {len(data):,} bytes, type {mime or 'unknown'}]\n"
                    f"First 64 bytes (hex): {hexhead}\n\nPrintable strings:\n{_strings(data)}"), ""


def save_upload(name: str, data: bytes, ocr=True):
    """Store an upload and what can be read from it. ocr: work out an image's text now (1.3), while this runs in a
    thread; screenshots pass False (their tools read them themselves, only for models without vision)."""
    uid = uuid.uuid4().hex[:16]
    safe = re.sub(r"[^\w.\- ]", "_", name)[:120] or "file"
    folder = UPLOADS / uid
    folder.mkdir(parents=True, exist_ok=True)
    (folder / safe).write_bytes(data)
    kind, text, note = extract(name, data)
    truncated = False
    if text and len(text) > MAX_CHARS:
        text, truncated = text[:MAX_CHARS], True
    if text is not None:
        (folder / "extracted.txt").write_text(text, encoding="utf-8")
    meta = {"id": uid, "name": name, "file": safe, "kind": kind, "size": len(data), "note": note,
            "chars": len(text) if text else 0, "truncated": truncated,
            "mime": mimetypes.guess_type(name)[0] or "application/octet-stream"}
    (folder / "meta.json").write_text(json.dumps(meta), encoding="utf-8")
    if ocr and kind == "image":
        try:
            ocr_text(uid)
        except Exception:  # noqa: BLE001
            pass
    return meta


def _on_event_loop():
    import asyncio
    try:
        asyncio.get_running_loop()
        return True
    except RuntimeError:
        return False


def ocr_text(uid: str):
    """The text in an attached image (OCR, worked out once and kept next to it); "" when there is none."""
    folder = UPLOADS / re.sub(r"\W", "", uid)
    cache = folder / "ocr.txt"
    if cache.exists():
        return cache.read_text(encoding="utf-8")
    from . import ocr
    p = file_path(uid)
    if not p or not p.exists() or not ocr.available() or _on_event_loop():
        return ""                     # never on the server's event loop: that would hold up every other request
    t = ocr.image_text(p.read_bytes())
    try:
        cache.write_text(t, encoding="utf-8")
    except OSError:
        pass
    return t


def load_meta(uid: str):
    p = UPLOADS / re.sub(r"\W", "", uid) / "meta.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def file_path(uid: str):
    m = load_meta(uid)
    return (UPLOADS / m["id"] / m["file"]) if m else None


def text_of(uid: str):
    p = UPLOADS / re.sub(r"\W", "", uid) / "extracted.txt"
    return p.read_text(encoding="utf-8") if p.exists() else ""


def image_data_url(uid: str, max_side=1568):
    """Downscale big images so they don't blow up the vision token budget."""
    from PIL import Image
    p = file_path(uid)
    im = Image.open(p)
    im.load()
    if im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    w, h = im.size
    s = max_side / max(w, h)
    if s < 1:
        im = im.resize((int(w * s), int(h * s)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def save_image_bytes(png: bytes, name="screenshot.png"):
    return save_upload(name, png, ocr=False)
